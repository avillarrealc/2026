import os, re
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, fill_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=120, bottom=120, left=160, right=160):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_cell_borders(cell, color='CBD5E1', sz='4', val='single'):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/><w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/><w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/><w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/></w:tcBorders>')
    tcPr.append(tcBorders)

base_dir = r"d:\Documentos\IEJAGA\2026\Grado 11\Clase 03 - Rodaje y Grabacion del Anuncio"
md_path = os.path.join(base_dir, "planeacion.md")
docx_path = os.path.join(base_dir, "planeacion.docx")

with open(md_path, 'r', encoding='utf-8') as f:
    text = f.read()

doc = docx.Document()

# Margins
for s in doc.sections:
    s.top_margin = Inches(0.8)
    s.bottom_margin = Inches(0.8)
    s.left_margin = Inches(0.9)
    s.right_margin = Inches(0.9)

PRIMARY = RGBColor(30, 58, 138)    # Navy Blue
SECONDARY = RGBColor(37, 99, 235)  # Blue
TEXT_COLOR = RGBColor(30, 41, 59)   # Slate
MUTED_COLOR = RGBColor(100, 116, 139)

# Institution Header
p_inst = doc.add_paragraph()
p_inst.paragraph_format.space_after = Pt(2)
r_inst = p_inst.add_run('INSTITUCIÓN EDUCATIVA JOSÉ ANTONIO GALÁN - LA ESTRELLA')
r_inst.bold = True
r_inst.font.name = 'Arial'
r_inst.font.size = Pt(11)
r_inst.font.color.rgb = PRIMARY

p_sub = doc.add_paragraph()
p_sub.paragraph_format.space_after = Pt(14)
r_sub = p_sub.add_run('Área de Tecnología e Informática · Planeación de Clase (90 Minutos)')
r_sub.font.name = 'Arial'
r_sub.font.size = Pt(9.5)
r_sub.font.color.rgb = MUTED_COLOR

lines = text.split('\n')
i = 0
in_table = False
table_rows = []

def process_table(rows):
    if not rows:
        return
    num_cols = max(len(r) for r in rows)
    tbl = doc.add_table(rows=len(rows), cols=num_cols)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = True
    
    for r_idx, row_data in enumerate(rows):
        is_header = (r_idx == 0)
        for c_idx, val in enumerate(row_data):
            if c_idx >= num_cols:
                continue
            cell = tbl.cell(r_idx, c_idx)
            cell.text = ''
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.1
            
            clean_val = val.replace('**', '').strip()
            run = p.add_run(clean_val)
            run.font.name = 'Arial'
            
            if is_header:
                set_cell_background(cell, '1E3A8A')
                run.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.font.size = Pt(9)
            else:
                bg = 'F8FAFC' if r_idx % 2 == 1 else 'FFFFFF'
                set_cell_background(cell, bg)
                run.font.size = Pt(8.5)
                run.font.color.rgb = TEXT_COLOR
                if val.strip().startswith('**') and val.strip().endswith('**'):
                    run.bold = True
                    
            set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
            set_cell_borders(cell, color='CBD5E1', sz='4')
            
    doc.add_paragraph().paragraph_format.space_after = Pt(8)

while i < len(lines):
    line = lines[i].strip()
    if not line:
        if in_table:
            process_table(table_rows)
            table_rows = []
            in_table = False
        i += 1
        continue
    
    if line.startswith('|') and line.endswith('|'):
        if '---' in line:
            i += 1
            continue
        cells = [c.strip() for c in line.split('|')[1:-1]]
        table_rows.append(cells)
        in_table = True
        i += 1
        continue
    else:
        if in_table:
            process_table(table_rows)
            table_rows = []
            in_table = False

    if line.startswith('# '):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(14)
        h.paragraph_format.space_after = Pt(6)
        r = h.add_run(line[2:].strip())
        r.bold = True
        r.font.name = 'Arial'
        r.font.size = Pt(13)
        r.font.color.rgb = PRIMARY
    elif line.startswith('### '):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(10)
        h.paragraph_format.space_after = Pt(4)
        r = h.add_run(line[4:].strip())
        r.bold = True
        r.font.name = 'Arial'
        r.font.size = Pt(11)
        r.font.color.rgb = SECONDARY
    elif line.startswith('- **'):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        parts = line[2:].split('**')
        bold_part = parts[1] if len(parts) > 1 else ''
        rest_part = parts[2] if len(parts) > 2 else ''
        rb = p.add_run(bold_part)
        rb.bold = True
        rb.font.name = 'Arial'
        rb.font.size = Pt(10)
        rb.font.color.rgb = TEXT_COLOR
        rr = p.add_run(rest_part)
        rr.font.name = 'Arial'
        rr.font.size = Pt(10)
        rr.font.color.rgb = TEXT_COLOR
    elif line.startswith('- '):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        tokens = re.split(r'(\*\*.*?\*\*)', line[2:])
        for t in tokens:
            if t.startswith('**') and t.endswith('**'):
                r = p.add_run(t[2:-2])
                r.bold = True
            else:
                r = p.add_run(t)
            r.font.name = 'Arial'
            r.font.size = Pt(10)
            r.font.color.rgb = TEXT_COLOR
    else:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.15
        tokens = re.split(r'(\*\*.*?\*\*)', line)
        for t in tokens:
            if t.startswith('**') and t.endswith('**'):
                r = p.add_run(t[2:-2])
                r.bold = True
            else:
                r = p.add_run(t)
            r.font.name = 'Arial'
            r.font.size = Pt(10)
            r.font.color.rgb = TEXT_COLOR
    i += 1

if in_table:
    process_table(table_rows)

doc.save(docx_path)
print("SUCCESS")
